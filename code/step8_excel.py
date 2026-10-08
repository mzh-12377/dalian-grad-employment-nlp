# -*- coding: utf-8 -*-
"""Step 8 · 汇总导出 Excel 结果簿（浏览结果表交付物）"""
import os
import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
M = os.path.join(ROOT, "data_measure")
OUT = os.path.join(ROOT, "results", "浏览结果表-全部结果.xlsx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)

sheets = {
    "岗位测度主表": os.path.join(M, "job_measures.csv"),
    "描述性统计": os.path.join(M, "descriptive.csv"),
    "信效度检验": os.path.join(M, "reliability.csv"),
    "维度相关矩阵": os.path.join(M, "validity_dim_corr.csv"),
    "效标效度": os.path.join(M, "validity_criterion.csv"),
    "主回归": os.path.join(M, "regression_main.csv"),
    "简约回归": os.path.join(M, "regression_parsimonious.csv"),
    "分行业回归": os.path.join(M, "regression_by_industry.csv"),
    "共线性诊断": os.path.join(M, "regression_vif.csv"),
    "高薪分类": os.path.join(M, "classification_highpay.csv"),
    "聚类画像": os.path.join(M, "cluster_profiles.csv"),
}

with pd.ExcelWriter(OUT, engine="openpyxl") as w:
    for name, path in sheets.items():
        if not os.path.exists(path):
            continue
        df = pd.read_csv(path)
        # 列数过多的主表裁剪关键列
        if name == "岗位测度主表":
            keep = [c for c in ["job_id","title","city","company","industry","comp_type","salary_raw",
                                "salary_mid","log_salary","DTS","GCS","PBS","PES","SKILL_TOTAL","EDU","GRF",
                                "welfare_n","skill_hits"] if c in df.columns]
            df = df[keep]
        df.to_excel(w, sheet_name=name[:31], index=False)

# 表头美化
from openpyxl import load_workbook
wb = load_workbook(OUT)
head_fill = PatternFill("solid", fgColor="1F4E79")
head_font = Font(color="FFFFFF", bold=True, name="Microsoft YaHei", size=10)
for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    for c in ws[1]:
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(horizontal="center", vertical="center")
    for col in ws.columns:
        width = max((len(str(c.value)) if c.value else 0) for c in col[:50])
        ws.column_dimensions[col[0].column_letter].width = min(28, max(10, width * 1.6))
wb.save(OUT)

print("Excel 已生成:", OUT, os.path.getsize(OUT), "字节")
print("工作表:", ", ".join(wb.sheetnames))
